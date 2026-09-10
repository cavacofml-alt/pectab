#!/usr/bin/env python3
"""Parse the "bagtag specs" xlsx (from whoever manages/creates PECTABs)
into data/pectabs.json — supersedes scripts/parse-pectabs.py's ADD.txt/
PAX.txt-derived catalog as the primary source: richer (extra fields like
"currently in use", free-text limiting notes) and more current (176 vs
127 records; the 5 that overlap with different values were confirmed
against real remarks in the sheet, e.g. P5801 was repurposed to a
no-stubs kiosk layout).

Usage:
    python3 scripts/parse-bagtag-specs-xlsx.py path/to/bagtag-specs.xlsx > data/pectabs.json

`dest` (destinations supported) has no column in this sheet but existed
in the old ADD.txt/PAX.txt-derived data/pectabs.json — pass that file
with --dest-source to carry it forward for ids present in both:

    python3 scripts/parse-bagtag-specs-xlsx.py specs.xlsx --dest-source data/pectabs.json > data/pectabs.json.new
"""
import argparse
import json
import sys

try:
    import openpyxl
except ImportError:
    print("needs openpyxl: pip install openpyxl", file=sys.stderr)
    raise

SHEET_NAME = "bagtag specs"
HEADER = [
    "Pectabnr", "orientation", "nr of additional stubs", "total tag length",
    "Length of main tagpart", "length of one additional stub",
    "additional stubs are all same size", "length of pax stub",
    "limiting features", "Currently in use", "limiting", "remarks", "Mirrorpoint",
]


def yn(value):
    if value is None:
        return None
    s = str(value).strip().upper()
    if s == "Y":
        return True
    if s == "N":
        return False
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--dest-source", help="old data/pectabs.json to carry the dest field forward from")
    args = ap.parse_args()

    dest_by_id = {}
    if args.dest_source:
        with open(args.dest_source, encoding="utf-8") as f:
            for rec in json.load(f):
                if "dest" in rec:
                    dest_by_id[rec["id"]] = rec["dest"]

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[SHEET_NAME]
    actual_header = [ws.cell(row=1, column=c).value for c in range(1, len(HEADER) + 1)]
    assert actual_header == HEADER, f"unexpected header: {actual_header}"

    records = []
    skipped = []
    for r in range(2, ws.max_row + 1):
        nr = ws.cell(row=r, column=1).value
        if nr is None:
            continue
        pid = f"P{int(nr):04d}"

        orientation = ws.cell(row=r, column=2).value
        dir_ = "PAX" if orientation == "pax stub first" else "ADD" if orientation == "add. stub first" else None

        st = ws.cell(row=r, column=3).value
        length = ws.cell(row=r, column=4).value
        main = ws.cell(row=r, column=5).value
        add = ws.cell(row=r, column=6).value
        eq = yn(ws.cell(row=r, column=7).value)
        pax = ws.cell(row=r, column=8).value
        in_use = yn(ws.cell(row=r, column=10).value)
        limiting = ws.cell(row=r, column=11).value
        remarks_col = ws.cell(row=r, column=12).value
        mirrorpoint = ws.cell(row=r, column=13).value
        note_n = ws.cell(row=r, column=14).value

        if dir_ is None or st is None or length is None or main is None or add is None or pax is None:
            # linhas incompletas (ex: calibração) — sem estrutura suficiente para o motor de matching
            skipped.append((pid, "incomplete core fields"))
            continue

        remarks_parts = [str(p).strip() for p in (limiting, remarks_col, note_n) if p not in (None, "")]
        rec = {
            "id": pid,
            "dir": dir_,
            "st": int(st),
            "len": int(length),
            "pax": int(pax),
            "main": int(main),
            "add": int(add),
        }
        if eq is not None:
            rec["eq"] = eq
        if in_use is not None:
            rec["inUse"] = in_use
        if pid in dest_by_id:
            rec["dest"] = dest_by_id[pid]
        if remarks_parts:
            rec["remarks"] = " — ".join(remarks_parts)
        if mirrorpoint is not None:
            rec["mirrorPoint"] = mirrorpoint

        records.append(rec)

    records.sort(key=lambda r: r["id"])
    print(json.dumps(records, indent=2, ensure_ascii=False))

    print(f"TOTAL: {len(records)} (skipped {len(skipped)} incomplete rows)", file=sys.stderr)
    not_in_use = [r["id"] for r in records if r.get("inUse") is False]
    print(f"NOTE: {len(not_in_use)} records marked NOT currently in use: {not_in_use}", file=sys.stderr)
    unknown_in_use = [r["id"] for r in records if "inUse" not in r]
    print(f"NOTE: {len(unknown_in_use)} records with unknown in-use status: {unknown_in_use}", file=sys.stderr)


if __name__ == "__main__":
    main()
